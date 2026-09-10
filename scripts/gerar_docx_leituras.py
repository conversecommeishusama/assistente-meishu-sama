#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera arquivos .docx (Word) formatados para impressão a partir dos arquivos
markdown da apostila de leituras integrais do ciclo
"Mundo Espiritual e Antepassados".

Cada arquivo .md vira um .docx com:
  - Título do encontro (grande)
  - Metadado da fonte (itálico, pequeno)
  - Títulos de capítulo destacados (negrito, cor escura)
  - Parágrafos justificados, fonte serifada, entrelinha confortável
  - Falas/diálogos com o rótulo em negrito (Meishu-Sama:, Interlocutor: etc.)
  - Cabeçalho com o nome do ciclo e rodapé com numeração de página
"""
import os
import re
import glob

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- paths
BASE = os.path.dirname(os.path.abspath(__file__))
SRC_DIR = os.path.join(BASE, "..", "docs", "leituras_integrais")
OUT_DIR = os.path.join(BASE, "..", "docs", "leituras_word")
os.makedirs(OUT_DIR, exist_ok=True)

NOME_CICLO = "Ciclo de Estudos — Mundo Espiritual e Antepassados"
NOME_RODAPE = "Preparação para o Culto às Almas dos Antepassados — Solo Sagrado de Guarapiranga (02/11)"

COR_TITULO = RGBColor(0x1F, 0x3A, 0x5F)   # azul-escuro
COR_CAPITULO = RGBColor(0x2E, 0x2E, 0x2E)  # quase preto
COR_FONTE_META = RGBColor(0x55, 0x55, 0x55)
FONTE_TEXTO = "Georgia"
FONTE_TITULO = "Georgia"
TAM_TEXTO = 12

# Rótulos de diálogo que devem aparecer em negrito no início do parágrafo
ROTULOS = [
    "Interlocutor:", "Meishu-Sama:", "Mestre:", "Pergunta:", "Resposta:",
    "Ele:", "Eu:", "M:",
]

# ---------------------------------------------------------------- helpers

def eh_separador(linha):
    """Linha composta apenas por caracteres decorativos."""
    s = linha.strip()
    if not s:
        return True
    return all(c in "─━-—=_*•·" for c in s)


def eh_titulo(linha):
    """Heurística: linha curta, sem pontuação final de frase e que não seja
    diálogo nem comece com rótulo de metadado."""
    s = linha.strip()
    if not s or len(s) > 80:
        return False
    # remove marcações markdown simples p/ avaliar
    t = re.sub(r"[*_#>`~]", "", s).strip()
    if not t:
        return False
    # termina com pontuação de frase -> não é título
    if t[-1] in ".!?;:,":
        return False
    # começa com rótulo de diálogo -> não é título
    if any(t.startswith(r) for r in ROTULOS):
        return False
    # começa com números de lista
    if re.match(r"^[\d\(\)\[\]\-–—•·*]+", t):
        return False
    return True


def limpar_markdown(texto):
    """Remove marcações markdown simples do texto."""
    texto = texto.replace("**", "")
    texto = re.sub(r"^#{1,6}\s*", "", texto)
    texto = re.sub(r"^>\s?", "", texto)
    texto = texto.replace("`", "")
    texto = re.sub(r"^\[/?[a-z]+\]", "", texto)
    return texto.strip()


def estilizar_paragrafo(par, texto, tamanho=TAM_TEXTO, justificar=True):
    """Adiciona texto a um parágrafo, aplicando negrito aos rótulos de diálogo."""
    # negrito para rótulo inicial (ex.: "Meishu-Sama:")
    prefixo = None
    corpo = texto
    for r in ROTULOS:
        if texto.startswith(r):
            prefixo = r
            corpo = texto[len(r):].lstrip()
            break
    run_prefixo = None
    if prefixo:
        run_prefixo = par.add_run(prefixo + " ")
        run_prefixo.bold = True
        run_prefixo.font.name = FONTE_TEXTO
        run_prefixo.font.size = Pt(tamanho)
    run = par.add_run(corpo)
    run.font.name = FONTE_TEXTO
    run.font.size = Pt(tamanho)
    if justificar:
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = par.paragraph_format
    pf.line_spacing = 1.3
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)
    return par


def add_numero_pagina(par):
    """Adiciona campo de número de página a um parágrafo de rodapé."""
    run = par.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
    run.font.name = FONTE_TEXTO
    run.font.size = Pt(10)


MAX_BLOCO = 850  # caracteres máximos por bloco de parágrafo (p/ leitura/impressão)


def quebrar_blocos(texto, limite=MAX_BLOCO):
    """Divide um parágrafo longo em blocos menores, quebrando em fronteiras de
    frase (após '. ', '! ', '? ') ou, na falta, em espaços. Preserva a ordem e
    o texto integral (apenas remove espaços duplicados na quebra)."""
    texto = texto.strip()
    if len(texto) <= limite:
        return [texto]
    blocos = []
    restante = texto
    while len(restante) > limite:
        fatia = restante[:limite]
        # procurar última fronteira de frase dentro da fatia
        corte = -1
        for m in re.finditer(r'[\.\!\?]\s+', fatia):
            corte = m.end()
        # se não achou fronteira de frase, usar último espaço
        if corte < 0:
            corte = fatia.rfind(' ')
        if corte <= 0:
            corte = limite
        bloco = restante[:corte].strip()
        blocos.append(bloco)
        restante = restante[corte:].strip()
    if restante:
        blocos.append(restante)
    return blocos


# ---------------------------------------------------------------- geração

def gerar_docx(arquivo_md, nome_saida):
    doc = Document()

    # ---- margens da página (A4) ----
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    # estilo Normal base
    style = doc.styles["Normal"]
    style.font.name = FONTE_TEXTO
    style.font.size = Pt(TAM_TEXTO)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE_TEXTO)

    # ---- cabeçalho ----
    header = sec.header
    hp = header.paragraphs[0]
    hp.text = NOME_CICLO
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in hp.runs:
        run.font.name = FONTE_TEXTO
        run.font.size = Pt(9)
        run.font.color.rgb = COR_FONTE_META
        run.italic = True
    # linha fina sob o cabeçalho
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1F3A5F')
    pBdr.append(bottom); pPr.append(pBdr)

    # ---- rodapé com número de página ----
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_numero_pagina(fp)

    # ---------------------------------------------------------- conteúdo
    with open(arquivo_md, encoding="utf-8") as f:
        linhas = f.read().splitlines()

    # primeiro título (## Ensinamento — ...)
    titulo_doc = None
    for i, linha in enumerate(linhas):
        if linha.strip().startswith("##"):
            titulo_doc = limpar_markdown(linha)
            # remove até o título
            linhas = linhas[i + 1:]
            break

    # ---- Título principal ----
    if titulo_doc:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(titulo_doc.replace("Ensinamento — ", "").strip())
        r.bold = True
        r.font.name = FONTE_TITULO
        r.font.size = Pt(20)
        r.font.color.rgb = COR_TITULO
        p.paragraph_format.space_after = Pt(4)

        # subtítulo do ciclo
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run("Encontro de leitura de ensinamento de Meishu-Sama")
        r2.italic = True
        r2.font.name = FONTE_TEXTO
        r2.font.size = Pt(11)
        r2.font.color.rgb = COR_FONTE_META
        p2.paragraph_format.space_after = Pt(14)

    # linha separadora
    psep = doc.add_paragraph()
    psep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rsep = psep.add_run("―" * 30)
    rsep.font.color.rgb = COR_TITULO
    rsep.font.size = Pt(10)
    psep.paragraph_format.space_after = Pt(12)

    primeiro_paragrafo = True
    for linha in linhas:
        if eh_separador(linha):
            continue
        s = linha.strip()

        # metadado da fonte (**Fonte:** / **Fonte N:**)
        m = re.match(r"^\*\*(Fonte[^*]*?):\*\*\s*(.*)$", s)
        if m:
            p = doc.add_paragraph()
            r1 = p.add_run(m.group(1) + ": ")
            r1.bold = True
            r1.font.name = FONTE_TEXTO
            r1.font.size = Pt(10.5)
            r1.font.color.rgb = COR_FONTE_META
            r2 = p.add_run(limpar_markdown(m.group(2)))
            r2.italic = True
            r2.font.name = FONTE_TEXTO
            r2.font.size = Pt(10.5)
            r2.font.color.rgb = COR_FONTE_META
            p.paragraph_format.space_after = Pt(10)
            continue

        # citação de apoio (**Leitura complementar** etc.)
        m = re.match(r"^\*\*(.*?):\*\*\s*(.*)$", s)
        if m and len(m.group(1)) < 45:
            p = doc.add_paragraph()
            r1 = p.add_run(m.group(1) + ": ")
            r1.bold = True
            r1.font.name = FONTE_TEXTO
            r1.font.size = Pt(11)
            r1.font.color.rgb = COR_CAPITULO
            r2 = p.add_run(limpar_markdown(m.group(2)))
            r2.font.name = FONTE_TEXTO
            r2.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(8)
            continue

        # linha de separação markdown "---"
        if s == "---":
            psep2 = doc.add_paragraph()
            psep2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = psep2.add_run("•  •  •")
            rr.font.color.rgb = COR_TITULO
            rr.font.size = Pt(11)
            psep2.paragraph_format.space_before = Pt(6)
            psep2.paragraph_format.space_after = Pt(6)
            continue

        # título de capítulo (curto)
        if eh_titulo(s):
            p = doc.add_paragraph()
            r = p.add_run(limpar_markdown(s))
            r.bold = True
            r.font.name = FONTE_TITULO
            r.font.size = Pt(15)
            r.font.color.rgb = COR_CAPITULO
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(8)
            continue

        # parágrafo normal
        texto = limpar_markdown(s)
        if texto:
            for bloco in quebrar_blocos(texto):
                p = doc.add_paragraph()
                estilizar_paragrafo(p, bloco)
            primeiro_paragrafo = False

    caminho = os.path.join(OUT_DIR, nome_saida)
    doc.save(caminho)
    print(f"OK -> {nome_saida}  (origem: {os.path.basename(arquivo_md)})")
    return caminho


# ---------------------------------------------------------------- executar
if __name__ == "__main__":
    mapa = {
        "Aula_01_Existencia_do_Mundo_Espiritual.md": "Encontro_01_A_Existencia_do_Mundo_Espiritual.docx",
        "Aula_02_Constituicao_do_Mundo_Espiritual.md": "Encontro_02_A_Constituicao_do_Mundo_Espiritual.docx",
        "Aula_03_Culto_às_Almas_Obon.md": "Encontro_03_O_Culto_as_Almas_dos_Antepassados_Obon.docx",
        "Aula_04_Elo_Espiritual_e_Destino.md": "Encontro_04_Elo_Espiritual_e_Destino.docx",
        "Aula_05_Morte_Julgamento_e_Destino.md": "Encontro_05_Morte_Julgamento_e_Destino.docx",
        "Aula_06_Vida_no_Mundo_Espiritual_e_Reencarnacao.md": "Encontro_06_Vida_no_Mundo_Espiritual_e_Reencarnacao.docx",
        "Aula_07_Influencia_dos_Antepassados.md": "Encontro_07_A_Influencia_dos_Antepassados.docx",
        "Aula_08_Culto_na_Pratica.md": "Encontro_08_O_Culto_na_Pratica_e_Preparacao_02_11.docx",
    }

    for md, saida in mapa.items():
        origem = os.path.join(SRC_DIR, md)
        if os.path.exists(origem):
            gerar_docx(origem, saida)
        else:
            print("!! arquivo não encontrado:", origem)

    print("\nDocumentos gerados em:", OUT_DIR)
